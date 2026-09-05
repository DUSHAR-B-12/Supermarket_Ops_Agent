import os
import logging
from typing import Dict, Any, Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from app.config.settings import settings

logger = logging.getLogger(__name__)

# ── GreenBasket Supermart Brand Palette (PPTX) ──────────────────────
BRAND_DARK_GREEN = RGBColor(27, 94, 32)     # #1B5E20 — titles, main brand
BRAND_GREEN = RGBColor(46, 125, 50)         # #2E7D32 — section accents
BRAND_LIGHT_GREEN = RGBColor(67, 160, 71)   # #43A047 — highlights
DARK_GRAY = RGBColor(45, 55, 72)            # body text


def generate_analysis_deck_pptx(daily_data: Dict[str, Any], output_dir: Optional[str] = None) -> str:
    """
    Generate a 6-slide PowerPoint (.pptx) presentation deck based on database sales data.
    """
    if not output_dir:
        base_dir = os.path.dirname(os.path.abspath(__file__))
        output_dir = os.path.join(base_dir, "decks")
    os.makedirs(output_dir, exist_ok=True)

    date_str = daily_data.get("date", "Today")
    filename = f"Sales_Analysis_{date_str}.pptx"
    file_path = os.path.join(output_dir, filename)

    prs = Presentation()

    shop_name = settings.SHOP_NAME

    def add_title(slide, text):
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.4), Inches(0.8))
        tf = tx_box.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = BRAND_DARK_GREEN

    blank_layout = prs.slide_layouts[6]

    # Slide 1: Title & Daily Business Summary
    slide1 = prs.slides.add_slide(blank_layout)
    add_title(slide1, f"Daily Business Summary — {date_str}")
    tb1 = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = f"🏪 {shop_name} — Operations Report for {date_str}"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRAND_GREEN

    bullets = [
        f"💰 Total Sales Revenue: ₹ {daily_data.get('total_sales', 0.0):.2f}",
        f"🧾 Total Bills Finalized: {daily_data.get('bill_count', 0)}",
        f"📊 Total Tax Collected (CGST+SGST): ₹ {daily_data.get('total_tax', 0.0):.2f}",
        f"📑 Net Subtotal Revenue: ₹ {daily_data.get('subtotal', 0.0):.2f}",
        f"💳 Outstanding Khata Credit: ₹ {daily_data.get('total_khata_outstanding', 0.0):.2f}",
    ]
    for b in bullets:
        p = tf1.add_paragraph()
        p.text = b
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(10)

    # Slide 2: Sales Breakdown by Payment Method
    slide2 = prs.slides.add_slide(blank_layout)
    add_title(slide2, "Sales Breakdown by Payment Method")
    tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    breakdown = daily_data.get("payment_breakdown", {})
    counts = daily_data.get("payment_counts", {})
    
    for method in ["Cash", "UPI", "Card", "Khata"]:
        val = breakdown.get(method, 0.0)
        cnt = counts.get(method, 0)
        p = tf2.add_paragraph() if tf2.paragraphs[0].text else tf2.paragraphs[0]
        p.text = f"• {method}: ₹ {val:.2f} ({cnt} transactions)"
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(12)

    # Slide 3: Top Products Sold
    slide3 = prs.slides.add_slide(blank_layout)
    add_title(slide3, "Top Products Sold Today")
    tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(4.5))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    top_prods = daily_data.get("top_selling_products", [])
    if top_prods:
        for idx, prod in enumerate(top_prods, 1):
            p = tf3.add_paragraph() if tf3.paragraphs[0].text else tf3.paragraphs[0]
            p.text = f"{idx}. {prod['name']} — {prod['quantity_sold']} {prod['unit']} (Revenue: ₹ {prod['revenue']:.2f})"
            p.font.size = Pt(15)
            p.font.color.rgb = DARK_GRAY
            p.space_after = Pt(10)
    else:
        p = tf3.paragraphs[0]
        p.text = "No items sold on this date."
        p.font.size = Pt(15)

    # Slide 4: Inventory Health & Reorder Alerts
    slide4 = prs.slides.add_slide(blank_layout)
    add_title(slide4, "Inventory Health & Low-Stock Alerts")
    tb4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(4.5))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    low_stock = daily_data.get("low_stock_products", [])
    if low_stock:
        p = tf4.paragraphs[0]
        p.text = f"⚠️ {len(low_stock)} Items Running Low On Stock:"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(197, 48, 48)
        p.space_after = Pt(10)

        for prod in low_stock[:6]:
            p = tf4.add_paragraph()
            p.text = f"• {prod['name']}: {prod['quantity']} {prod['unit']} left (Reorder Level: {prod['reorder_level']})"
            p.font.size = Pt(14)
            p.font.color.rgb = DARK_GRAY
            p.space_after = Pt(6)
    else:
        p = tf4.paragraphs[0]
        p.text = "✅ Stock levels are healthy. No items below reorder levels."
        p.font.size = Pt(16)
        p.font.color.rgb = BRAND_LIGHT_GREEN

    # Slide 5: Khata Outstanding Credit
    slide5 = prs.slides.add_slide(blank_layout)
    add_title(slide5, "Khata Credit Ledger Summary")
    tb5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(4.5))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    total_khata = daily_data.get("total_khata_outstanding", 0.0)
    p = tf5.paragraphs[0]
    p.text = f"📋 Total Customer Credit Outstanding: ₹ {total_khata:.2f}"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRAND_GREEN
    p.space_after = Pt(15)

    khata_sales = daily_data.get("payment_breakdown", {}).get("Khata", 0.0)
    p = tf5.add_paragraph()
    p.text = f"• Today's Khata Credit Sales: ₹ {khata_sales:.2f}"
    p.font.size = Pt(15)
    p.font.color.rgb = DARK_GRAY

    # Slide 6: Business Insights & Recommendations
    slide6 = prs.slides.add_slide(blank_layout)
    add_title(slide6, "Operational Business Insights")
    tb6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(4.5))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    insights = [
        f"• Average Order Value: ₹ {(daily_data.get('total_sales', 0.0) / max(daily_data.get('bill_count', 1), 1)):.2f} per bill.",
        f"• Tax Compliance: Fully accounted ₹ {daily_data.get('total_cgst', 0.0):.2f} CGST + ₹ {daily_data.get('total_sgst', 0.0):.2f} SGST.",
        f"• Inventory Action: Reorder {len(low_stock)} low-stock SKUs to maintain availability.",
    ]
    for ins in insights:
        p = tf6.add_paragraph() if tf6.paragraphs[0].text else tf6.paragraphs[0]
        p.text = ins
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(12)

    prs.save(file_path)
    logger.info(f"Generated PPTX analysis deck at: {file_path}")
    return file_path
