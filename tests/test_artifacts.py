import os
import pytest
from unittest.mock import AsyncMock, MagicMock
from pptx import Presentation

from app.db.seed import seed_database
from app.db.models import Product, BillStatus
from app.services.billing_service import BillingService
from app.services.daily_close_service import DailyCloseService
from app.services.khata_service import KhataService
from app.artifacts.invoice import generate_pdf_invoice
from app.artifacts.analysis_deck import generate_analysis_deck_pptx
from app.tools.report_tools import generate_invoice_pdf, get_daily_close, generate_analysis_deck
from app.agent.registry import TOOL_MAP, TOOL_SCHEMAS, execute_tool
from app.telegram.bot import handle_message

# 1. Invoice Generator Creates Valid PDF
def test_invoice_pdf_creation(db_session, tmp_path):
    seed_database(db_session)
    atta = db_session.query(Product).filter(Product.sku == "SKU-ATTA-5KG").first()

    b_service = BillingService(db_session)
    bill = b_service.create_draft_bill()
    b_service.add_bill_item(bill.id, atta.id, 2.0)
    finalized = b_service.finalize_bill(bill.id, payment_method="Cash")

    pdf_path = generate_pdf_invoice(finalized, output_dir=str(tmp_path))
    assert os.path.exists(pdf_path)
    assert pdf_path.endswith(".pdf")

# 2. PDF Rejects Unfinalized Draft Bill
def test_pdf_rejects_unfinalized_bill(db_session, tmp_path):
    seed_database(db_session)
    b_service = BillingService(db_session)
    bill = b_service.create_draft_bill()

    with pytest.raises(ValueError, match="Must be 'finalized'"):
        generate_pdf_invoice(bill, output_dir=str(tmp_path))

# 3. Invoice PDF Filename & Number Check
def test_invoice_number_in_pdf(db_session, tmp_path):
    seed_database(db_session)
    atta = db_session.query(Product).filter(Product.sku == "SKU-ATTA-5KG").first()

    b_service = BillingService(db_session)
    bill = b_service.create_draft_bill()
    b_service.add_bill_item(bill.id, atta.id, 1.0)
    finalized = b_service.finalize_bill(bill.id, payment_method="UPI")

    pdf_path = generate_pdf_invoice(finalized, output_dir=str(tmp_path))
    assert finalized.bill_number in pdf_path

# 4. Invoice Uses Persisted Bill Totals
def test_invoice_persisted_totals(db_session, tmp_path):
    seed_database(db_session)
    atta = db_session.query(Product).filter(Product.sku == "SKU-ATTA-5KG").first()

    b_service = BillingService(db_session)
    bill = b_service.create_draft_bill()
    b_service.add_bill_item(bill.id, atta.id, 1.0)  # 240 + 5% = 252.0
    finalized = b_service.finalize_bill(bill.id, payment_method="Cash")

    assert float(finalized.subtotal) == 240.0
    assert float(finalized.total_tax) == 12.0
    assert float(finalized.grand_total) == 252.0

# 5-11. Daily Close Service Metrics Tests
def test_daily_close_service_metrics(db_session):
    seed_database(db_session)
    atta = db_session.query(Product).filter(Product.sku == "SKU-ATTA-5KG").first()
    salt = db_session.query(Product).filter(Product.sku == "SKU-SALT-1KG").first()

    b_service = BillingService(db_session)

    # Bill 1: Cash
    bill1 = b_service.create_draft_bill()
    b_service.add_bill_item(bill1.id, atta.id, 1.0)
    b_service.finalize_bill(bill1.id, payment_method="Cash")

    # Bill 2: UPI
    bill2 = b_service.create_draft_bill()
    b_service.add_bill_item(bill2.id, salt.id, 2.0)
    b_service.finalize_bill(bill2.id, payment_method="UPI")

    # Khata credit setup
    khata_service = KhataService(db_session)
    cust = khata_service.get_or_create_customer("Dinesh")
    khata_service.record_credit(cust.id, 450.0)

    daily_service = DailyCloseService(db_session)
    summary = daily_service.get_daily_summary()

    # 5. Total Sales
    assert summary["total_sales"] == 306.6  # 252.0 (Atta) + 54.6 (Salt)
    # 6. Bill Count
    assert summary["bill_count"] == 2
    # 7. Payment Method Breakdown
    assert summary["payment_breakdown"]["Cash"] == 252.0
    assert summary["payment_breakdown"]["UPI"] == 54.6
    # 8. GST Totals
    assert summary["total_tax"] == 14.6  # 12.0 + 2.6

    # 9. Top Selling Products
    assert len(summary["top_selling_products"]) >= 1
    # 10. Low Stock Products
    assert isinstance(summary["low_stock_products"], list)
    # 11. Khata Outstanding
    assert summary["total_khata_outstanding"] == 450.0

# 12-14. PPTX Analysis Deck Tests
def test_pptx_deck_generation(db_session, tmp_path):
    seed_database(db_session)
    daily_service = DailyCloseService(db_session)
    summary = daily_service.get_daily_summary()

    pptx_path = generate_analysis_deck_pptx(summary, output_dir=str(tmp_path))
    
    # 12. Deck File Created
    assert os.path.exists(pptx_path)
    assert pptx_path.endswith(".pptx")

    # 13. Deck Contains 6 Slides
    prs = Presentation(pptx_path)
    assert len(prs.slides) == 6

    # 14. Deck Contains Actual Sales Information
    text_runs = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text_runs.append(shape.text_frame.text)
    combined_text = " ".join(text_runs)
    assert "Daily Business Summary" in combined_text
    assert "Khata Credit Ledger Summary" in combined_text

# 15-17. Tool Registration Tests
def test_new_tools_registration():
    # 15. Invoice Tool Registered
    assert "generate_invoice_pdf" in TOOL_MAP
    # 16. Daily Close Tool Registered
    assert "get_daily_close" in TOOL_MAP
    # 17. Analysis Deck Tool Registered
    assert "generate_analysis_deck" in TOOL_MAP

# 18-20. Agent Tool Invocation Tests
def test_agent_tool_invocations(db_session, tmp_path):
    seed_database(db_session)
    atta = db_session.query(Product).filter(Product.sku == "SKU-ATTA-5KG").first()

    b_service = BillingService(db_session)
    bill = b_service.create_draft_bill()
    b_service.add_bill_item(bill.id, atta.id, 1.0)
    b_service.finalize_bill(bill.id, payment_method="Cash")

    # 18. Invoice tool invocation
    inv_res = execute_tool(db_session, "generate_invoice_pdf", {"bill_id": bill.id})
    assert inv_res["status"] == "success"
    assert "pdf_path" in inv_res["data"]

    # 19. Daily close tool invocation
    dc_res = execute_tool(db_session, "get_daily_close", {})
    assert dc_res["status"] == "success"
    assert "total_sales" in dc_res["data"]

    # 20. Analysis deck tool invocation
    deck_res = execute_tool(db_session, "generate_analysis_deck", {})
    assert deck_res["status"] == "success"
    assert "pptx_path" in deck_res["data"]

# 21-22. Telegram Document Delivery Tests
@pytest.mark.asyncio
async def test_telegram_document_delivery(tmp_path):
    # Create test dummy PDF and PPTX
    dummy_pdf = os.path.join(tmp_path, "test_invoice.pdf")
    with open(dummy_pdf, "w") as f:
        f.write("PDF Content")

    dummy_pptx = os.path.join(tmp_path, "test_deck.pptx")
    with open(dummy_pptx, "w") as f:
        f.write("PPTX Content")

    update = MagicMock()
    update.effective_user.id = 9999
    update.effective_user.username = "test_owner"
    update.message.text = "Send me the invoice"
    update.message.reply_text = AsyncMock()
    update.message.reply_document = AsyncMock()

    context = MagicMock()

    # Mock runner output containing [FILE: ...] tags
    from unittest.mock import patch
    with patch("app.telegram.bot.process_agent_message", new_callable=AsyncMock) as mock_agent:
        mock_agent.return_value = f"Here is your invoice:\n[FILE: {dummy_pdf}]\n[FILE: {dummy_pptx}]"
        
        await handle_message(update, context)

        # 21 & 22. Verification of document delivery
        assert update.message.reply_document.call_count == 2
